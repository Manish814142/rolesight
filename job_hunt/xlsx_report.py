"""A formatted Excel report (.xlsx) beside each CSV: the same rows, easier to read and to filter (T128).

The CSV stays the primary output (scripts, tests and the MCP tool read it). The report is built *from* the CSV that was just written, so
its columns, their order and the optional columns always agree with it, and it is the secondary output: a report that cannot be written
(the file is open in Excel on Windows, a missing library) is logged and never loses the scan.
"""
from __future__ import annotations

import csv
import logging
import os
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)

STRONG_SCORE = 85                      # green at or above this (the README's own example is `export --min 85`); amber from min_score up to it
HEADER_FILL = "1F3864"                 # dark blue header, white bold text
STRONG_COLOURS = ("C6EFCE", "006100")  # fill, text: Excel's "Good" style
WORTH_COLOURS = ("FFEB9C", "9C5700")   # Excel's "Neutral" style
LINK_COLOUR = "0563C1"
FONT, SIZE = "Calibri", 11              # named on every font set here: a font without a name is the viewer's choice (a serif one in LibreOffice)
LINE_HEIGHT, MAX_LINES = 15, 12         # points per line of 11-point text; a row is never taller than this many lines
LINK_LABELS = {"Application URL": "Open", "LinkedIn people": "Search"}   # the address is kept in the link; the cell stays short
MAX_LINK = 2000                        # Excel refuses a link address much beyond 2,079 characters
MAX_TEXT = 32000                       # a cell holds at most 32,767
WRAPPED = ("Role", "Location", "Reason")
WIDTHS = {"Score (%)": 10, "Company": 24, "ATS": 16, "Role": 44, "Location": 30, "Posted": 18, "Reason": 80, "Application URL": 12,
          "LinkedIn people": 12, "Worth Applying": 15, "Scan Date": 12}
DEFAULT_WIDTH = 16


def _number(text: str) -> Optional[int | float]:
    try:
        return int(text)
    except ValueError:
        try:
            return float(text)
        except ValueError:
            return None


def _font(**style):
    from openpyxl.styles import Font
    return Font(name=FONT, size=SIZE, **style)


def _lines(text: str, width: float) -> int:
    """Lines `text` needs in a wrapped column `width` characters wide (generous: Calibri fits a little more than the width)."""
    return sum(max(1, -(-len(part) // int(width))) for part in text.split("\n"))


def _build(header: list[str], rows: list[list[str]]):
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties

    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    wb.properties.title = "RoleSight job report"
    ws.append(header)
    head_font, head_fill = _font(bold=True, color="FFFFFF"), PatternFill("solid", start_color=HEADER_FILL, end_color=HEADER_FILL)
    for cell in ws[1]:
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 24

    link_font = _font(color=LINK_COLOUR, underline="single")
    widths = [WIDTHS.get(name, DEFAULT_WIDTH) for name in header]
    for r, row in enumerate(rows, start=2):
        lines = max([1, *(_lines(text, widths[c]) for c, (name, text) in enumerate(zip(header, row)) if name in WRAPPED and text)])
        if lines > 1:
            ws.row_dimensions[r].height = LINE_HEIGHT * min(lines, MAX_LINES)    # LibreOffice drew wrapped rows over each other without it (Excel not tested)
        for c, (name, text) in enumerate(zip(header, row), start=1):
            cell = ws.cell(r, c)
            cell.alignment = Alignment(vertical="top", wrap_text=name in WRAPPED)
            if not text:
                continue
            text = ILLEGAL_CHARACTERS_RE.sub("", text)[:MAX_TEXT]
            if name in LINK_LABELS and len(text) <= MAX_LINK:
                cell.value = LINK_LABELS[name]
                cell.hyperlink = text
                cell.font = link_font
            elif name == "Score (%)" and _number(text) is not None:
                cell.value = _number(text)
                cell.number_format = "0"
                cell.font = _font(bold=True)
                cell.alignment = Alignment(vertical="top", horizontal="center")
            else:
                cell.value = text
                cell.data_type = "s"       # text from a job board, never a formula, even when it begins with = + - @

    for c, name in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    last = get_column_letter(len(header))
    ws.freeze_panes = "C2"                 # the header row, and Score and Company, stay in view when scrolling
    ws.auto_filter.ref = f"A1:{last}{len(rows) + 1}"

    if rows and "Score (%)" in header:
        score = get_column_letter(header.index("Score (%)") + 1)
        area = f"{score}2:{score}{len(rows) + 1}"
        worth = get_column_letter(header.index("Worth Applying") + 1) if "Worth Applying" in header else None
        strong = f"${score}2>={STRONG_SCORE}"
        ws.conditional_formatting.add(area, FormulaRule(
            formula=[f'AND({strong},${worth}2<>"No")' if worth else strong], stopIfTrue=False,    # not green when the person's min_score rules it out
            fill=PatternFill("solid", start_color=STRONG_COLOURS[0], end_color=STRONG_COLOURS[0]), font=_font(bold=True, color=STRONG_COLOURS[1])))
        if worth:
            ws.conditional_formatting.add(area, FormulaRule(
                formula=[f'AND(${score}2<{STRONG_SCORE},${worth}2="Yes")'], stopIfTrue=False,
                fill=PatternFill("solid", start_color=WORTH_COLOURS[0], end_color=WORTH_COLOURS[0]), font=_font(bold=True, color=WORTH_COLOURS[1])))

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "1:1"
    return wb


def write_xlsx_report(csv_path: Path) -> Optional[Path]:
    """Write `csv_path` with `.xlsx` as a formatted workbook next to it; the path, or None when it could not be written (logged, and an
    older report of that name is removed so that it cannot be mistaken for this one). Never raises."""
    csv_path = Path(csv_path)
    target = csv_path.with_suffix(".xlsx")
    tmp = target.with_name(target.name + ".tmp")
    try:
        with csv_path.open(newline="", encoding="utf-8") as fh:
            header, *rows = csv.reader(fh)
        _build(header, rows).save(tmp)
        os.replace(tmp, target)            # never a half-written file if the old one is open
        return target
    except Exception as exc:               # noqa: BLE001 - the secondary output must not stop a scan
        logger.warning(f"Excel report for {csv_path.name} not written: {exc!r}. The CSV is complete.")
        for leftover in (tmp, target):
            try:
                leftover.unlink(missing_ok=True)
            except OSError:
                pass
        return None
