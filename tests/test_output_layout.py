"""T130: a clean output/. The two files you open every day keep fixed names at the top, the dated files go to history/, drafts to drafts/.

Manish (2026-10-17/18): "I want the output path clean so a user is not confused by so many files; the newest file should not be at the bottom
of 100 files." Chosen: the lighter variant (the CSV twins of the two fixed files stay at the top), "no monthly subfolders, leave my existing
files alone, no auto-delete". No network: adapters and scoring are stubbed; run_scan, the writers, the exporter and the SQLite rows are real.
"""
import csv
import logging
import os
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from job_hunt import drafter, scanner, tools
from job_hunt import main as cli
from job_hunt.db import init_db, update_job_status
from job_hunt.models import JobPosting, job_key

TODAY = date.today().isoformat()
TOP = {"history", "last_scan.csv", "last_scan.xlsx", "open_matches.csv", "open_matches.xlsx"}


def _job(n: int, ats: str = "greenhouse") -> JobPosting:
    return JobPosting(global_id=f"{ats}:x:{n}", ats=ats, company=f"Co {n}", title=f"Backend Engineer {n}", location="Bengaluru, India",
                      is_remote=False, url=f"https://x.example/{n}", description_text="Node.js", posted_at=TODAY)


@pytest.fixture
def scan(tmp_path, monkeypatch):
    """`scan(1, 2, 3, scores={"3": 40})`: one scan whose board lists jobs 1, 2 and 3; every job scores 80 unless `scores` says otherwise."""
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(scanner, "send_telegram", lambda *a, **k: False)
    (tmp_path / "resume.md").write_text("Node.js backend engineer")
    init_db("data/jobs.db")
    board: dict = {"jobs": [], "scores": {}}
    monkeypatch.setattr(scanner, "fetch_company_jobs", lambda c: list(board["jobs"]))

    def fake_score(job, **kw):
        score = board["scores"].get(job.title.rsplit(" ", 1)[1], 80)
        update_job_status(job_key(job), "scored", score, "fits")
        return {"score": score, "reason": "fits", "status": "scored"}

    monkeypatch.setattr(scanner, "score_and_process_job", fake_score)

    def run(*numbers, scores=None):
        board["jobs"], board["scores"] = [_job(n) for n in numbers], scores or {}
        scanner.run_scan({"candidate": {"resume_path": "resume.md", "min_score": 70}, "gemini_api_key": "k", "telegram": {}},
                         [{"name": "G", "ats": "greenhouse", "slug": "g"}])
    return run


def _names(path) -> set:
    return {p.name for p in Path(path).iterdir()}


def _rows(path) -> list[dict]:
    with Path(path).open(newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def _urls(path) -> list[str]:
    return [r["Application URL"].rsplit("/", 1)[1] for r in _rows(path)]


# --- the first scan in a fresh folder --------------------------------------------------------

def test_a_first_scan_in_a_fresh_folder_makes_exactly_the_agreed_layout(scan):
    scan(1, 2)
    assert _names("output") == TOP
    assert _names("output/history") == {f"jobs_{TODAY}.csv", f"jobs_{TODAY}.xlsx"}


def test_last_scan_holds_everything_the_scan_scored_best_first_and_open_matches_only_what_passes_min_score(scan):
    scan(1, 2, 3, scores={"1": 90, "2": 75, "3": 40})
    assert _urls("output/last_scan.csv") == ["1", "2", "3"]           # the 40 is here, below min_score 70
    assert _urls("output/open_matches.csv") == ["1", "2"]


def test_last_scan_is_the_same_table_as_the_dated_file_so_row_n_is_job_n_for_draft_and_mark(scan):
    scan(1, 2, 3, scores={"1": 90, "2": 75, "3": 40})
    dated = Path(f"output/history/jobs_{TODAY}.csv")
    assert Path("output/last_scan.csv").read_bytes() == dated.read_bytes()
    from job_hunt.db import last_scan_jobs
    assert [r["url"].rsplit("/", 1)[1] for r in last_scan_jobs()] == _urls("output/last_scan.csv")        # what `draft N` and `mark N` read


def test_both_reports_have_an_excel_copy_with_the_same_rows(scan):
    scan(1, 2, 3, scores={"3": 40})
    for name, csv_rows in (("last_scan", 3), ("open_matches", 2)):
        ws = openpyxl.load_workbook(f"output/{name}.xlsx").active
        assert ws.max_row - 1 == csv_rows == len(_rows(f"output/{name}.csv"))
    ws = openpyxl.load_workbook(f"output/history/jobs_{TODAY}.xlsx").active
    assert ws.max_row - 1 == 3


# --- later scans --------------------------------------------------------------------------------

def test_a_second_scan_the_same_day_adds_a_numbered_dated_file_and_replaces_the_fixed_ones(scan):
    scan(1, 2)
    scan(1, 2, 3, scores={"3": 95})
    assert _names("output/history") == {f"jobs_{TODAY}.csv", f"jobs_{TODAY}.xlsx", f"jobs_{TODAY}_1.csv", f"jobs_{TODAY}_1.xlsx"}
    assert _urls("output/last_scan.csv") == ["3"] == _urls("output/open_matches.csv")        # only what the last scan scored
    assert _urls(f"output/history/jobs_{TODAY}.csv") == ["1", "2"]                               # the first scan's file is never touched


def test_a_scan_that_scores_nothing_adds_no_dated_file_and_leaves_the_fixed_ones_empty_not_stale(scan):
    scan(1, 2)
    scan(1, 2)                                                           # nothing new
    assert _names("output/history") == {f"jobs_{TODAY}.csv", f"jobs_{TODAY}.xlsx"}
    for name in ("last_scan", "open_matches"):
        assert _rows(f"output/{name}.csv") == []
        ws = openpyxl.load_workbook(f"output/{name}.xlsx").active
        assert ws.max_row == 1 and ws.cell(1, 1).value == "Score (%)"       # a valid report with only the header


def test_the_very_first_scan_in_a_fresh_folder_that_scores_nothing_still_makes_the_four_fixed_files_and_no_history(scan):
    assert not Path("output").exists()
    scan()                                                               # the board lists no job at all
    assert _names("output") == {"last_scan.csv", "last_scan.xlsx", "open_matches.csv", "open_matches.xlsx"}
    assert _rows("output/last_scan.csv") == [] and _rows("output/open_matches.csv") == []


def test_a_scan_that_finds_only_weak_jobs_fills_last_scan_and_leaves_open_matches_empty(scan):
    scan(1, 2, scores={"1": 50, "2": 60})
    assert _urls("output/last_scan.csv") == ["2", "1"] and _rows("output/open_matches.csv") == []


def test_no_temporary_file_is_left_in_output_after_any_scan(scan):
    scan(1, 2)
    scan(1, 2, 3)
    assert not [p for p in Path("output").rglob("*") if p.name.endswith(".tmp")]


# --- what was there before is left alone -----------------------------------------------------------

def test_files_written_before_the_change_are_left_exactly_where_and_as_they_are(scan):
    Path("output").mkdir()
    old = {"jobs_2026-10-03.csv": b"old scan", "jobs_2026-10-04_1.csv": b"another", "ats_platform_verification.xlsx": b"xx", "open_matches.csv": b"old list"}
    for name, content in old.items():
        (Path("output") / name).write_bytes(content)
    scan(1, 2)
    for name, content in old.items():
        if name != "open_matches.csv":
            assert (Path("output") / name).read_bytes() == content          # not moved, not deleted, not rewritten
    assert _names("output") == TOP | {"jobs_2026-10-03.csv", "jobs_2026-10-04_1.csv", "ats_platform_verification.xlsx"}
    assert _urls("output/open_matches.csv") == ["1", "2"]                    # the fixed file itself is replaced, as before


def test_no_dated_file_is_ever_written_at_the_top_of_output_any_more(scan):
    scan(1)
    scan(1, 2)
    assert not list(Path("output").glob("jobs_*"))


# --- a fixed file that cannot be replaced (open in a spreadsheet on Windows) ----------------------------

def test_a_fixed_file_that_cannot_be_replaced_is_a_warning_not_the_end_of_the_scan(scan, monkeypatch, caplog):
    scan(1, 2)
    real = os.replace

    def locked(src, dst, *a, **k):
        if Path(dst).name in ("last_scan.csv", "open_matches.csv"):
            raise PermissionError(13, "in use", str(dst))
        return real(src, dst, *a, **k)

    monkeypatch.setattr(os, "replace", locked)
    caplog.clear()                                                       # the first scan's INFO lines also say "history"
    with caplog.at_level(logging.WARNING):
        scan(1, 2, 3)                                                    # must not raise
    warnings = [r.getMessage() for r in caplog.records if r.levelno >= logging.WARNING]
    assert len(warnings) == 2
    assert "last_scan.csv" in warnings[0] and "open_matches.csv" in warnings[1]
    assert all("output/history" in w and "spreadsheet" in w for w in warnings)         # what to do and where the results are
    assert _urls(f"output/history/jobs_{TODAY}_1.csv") == ["3"]          # the scan's own record is complete
    assert not [p for p in Path("output").rglob("*") if p.name.endswith(".tmp")]
    from job_hunt.db import last_scan_jobs
    assert [r["url"].rsplit("/", 1)[1] for r in last_scan_jobs()] == ["3"]   # the scan row was finished: `draft` and `mark` work


def test_the_exporters_own_default_folder_is_history_so_no_caller_can_put_a_dated_file_at_the_top(tmp_path, monkeypatch):
    from job_hunt.exporter import export_jobs_to_csv, get_incremental_csv_path
    monkeypatch.chdir(tmp_path)
    assert export_jobs_to_csv([{"company": "A", "title": "T", "score": 80}], base_name="x") == Path("output/history/x.csv")
    assert get_incremental_csv_path("y") == Path("output/history/y.csv")


# --- autopilot export ----------------------------------------------------------------------------------------

def test_autopilot_export_writes_into_history_prints_where_and_leaves_the_fixed_files_alone(scan, monkeypatch, capsys):
    scan(1, 2, 3, scores={"3": 40})
    before = {n: (Path("output") / n).read_bytes() for n in ("last_scan.csv", "open_matches.csv")}
    monkeypatch.setattr(sys, "argv", ["autopilot", "export", "--min", "70"])
    cli.main()
    out = capsys.readouterr().out
    assert f"output/history/jobs_{TODAY}_1.csv" in out and ".xlsx" in out
    assert _names("output/history") >= {f"jobs_{TODAY}_1.csv", f"jobs_{TODAY}_1.xlsx"}
    assert _names("output") == TOP
    assert {n: (Path("output") / n).read_bytes() for n in before} == before
    assert _urls(f"output/history/jobs_{TODAY}_1.csv") == ["1", "2"]


def test_autopilot_export_days_names_the_period_and_a_repeat_is_numbered(scan, monkeypatch):
    scan(1)
    monkeypatch.setattr(sys, "argv", ["autopilot", "export", "--days", "7"])
    cli.main()
    cli.main()
    assert {f"jobs_{TODAY}_last7d.csv", f"jobs_{TODAY}_last7d_1.csv"} <= _names("output/history")


def test_autopilot_export_with_nothing_to_export_writes_no_file(scan, monkeypatch, capsys):
    scan(1, scores={"1": 40})
    monkeypatch.setattr(sys, "argv", ["autopilot", "export", "--min", "99"])
    cli.main()
    assert "No jobs with score >= 99" in capsys.readouterr().out
    assert _names("output/history") == {f"jobs_{TODAY}.csv", f"jobs_{TODAY}.xlsx"}


# --- drafts and the MCP tools ----------------------------------------------------------------------------------------

def test_a_draft_goes_to_output_drafts_not_the_top_of_output(scan, monkeypatch):
    scan(1)
    monkeypatch.setattr(drafter, "chat_with_llm", lambda config, messages, **k: "LLM OUTPUT")
    (Path("resume")).mkdir()
    Path("resume/r.md").write_text("Ada, backend engineer")
    out = drafter.draft_application({"candidate": {"name": "Ada", "resume_path": "resume/r.md"}}, "1")
    assert Path("output/drafts").is_dir() and _names("output") == TOP | {"drafts"}
    (folder,) = Path("output/drafts").iterdir()
    assert folder.name == f"co_1-{date.today().isoformat()}" and "application_info.txt" in _names(folder)
    assert out == folder


def test_the_export_tool_says_which_file_it_wrote(scan, monkeypatch, tmp_path):
    scan(1, 2)
    (tmp_path / "config.json").write_text("{}")
    message = tools.tool_export(min_score=0, days=0, config_path=str(tmp_path / "config.json"))
    assert message.startswith("Export complete") and f"output/history/jobs_{TODAY}_1.csv" in message.replace("\\", "/")


def test_the_draft_tool_says_which_folder_it_wrote(scan, monkeypatch, tmp_path):
    monkeypatch.setattr(tools, "load_config", lambda: {})
    monkeypatch.setattr(tools, "draft_application", lambda cfg, ref: Path("output/drafts/co_1-2026-10-18"))
    (tmp_path / "config.json").write_text("{}")
    message = tools.tool_draft("#1", config_path=str(tmp_path / "config.json"))
    assert "output/drafts/co_1-2026-10-18" in message.replace("\\", "/")


# --- the guides --------------------------------------------------------------------------------------------------------

GUIDES = ["README.md", "CLAUDE.md", "ARCHITECTURE.md", "SETUP.md", "docs/00-user-guide.md", "docs/01-install.md", "docs/04-companies-and-scanning.md",
          "docs/05-outputs-and-notifications.md", "docs/06-mcp-and-skill.md", "docs/07-config-and-scoring.md", "docs/09-testing-checklist.md",
          "skills/rolesight/SKILL.md"]


def test_no_guide_sends_the_reader_to_a_dated_file_or_a_draft_folder_at_the_top_of_output():
    for name in GUIDES:
        text = (Path(__file__).resolve().parent.parent / name).read_text(encoding="utf-8")
        assert not re.search(r"output/jobs_", text), f"{name} still names output/jobs_… (it is output/history/jobs_…)"
        assert "output/<company>" not in text, f"{name} still names output/<company>-<date> (it is output/drafts/…)"


def test_the_guides_describe_the_new_files():
    root = Path(__file__).resolve().parent.parent
    for name in ("README.md", "docs/05-outputs-and-notifications.md", "docs/00-user-guide.md"):
        text = (root / name).read_text(encoding="utf-8")
        assert "last_scan" in text and "history/" in text, name
    assert "drafts/" in (root / "docs/05-outputs-and-notifications.md").read_text(encoding="utf-8")


def test_the_guides_say_how_a_link_opens_in_each_program():
    """Manish: without this a user on another system thinks the links are broken (LibreOffice asks for Ctrl+click by default)."""
    root = Path(__file__).resolve().parent.parent
    for name in ("docs/00-user-guide.md", "docs/05-outputs-and-notifications.md"):
        text = (root / name).read_text(encoding="utf-8")
        assert "Ctrl" in text and "LibreOffice" in text and "Excel" in text and "Google Sheets" in text, name
