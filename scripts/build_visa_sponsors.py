"""Which companies on the list appear in public visa-sponsor sources? (T133, groundwork for T108)

    python scripts/build_visa_sponsors.py                       # download the sources, write job_hunt/data/visa_sponsors.json
    python scripts/build_visa_sponsors.py --offline             # rebuild from the files downloaded last time (no network)
    python scripts/build_visa_sponsors.py --offline --us-h1b "$HOME/Downloads/Employer Information.csv"
                                                                # a newer H-1B hub export saved in a browser (T143)
    python scripts/build_visa_sponsors.py --report /tmp/m.csv   # also every match with the source's own name, to read by hand
    python scripts/build_visa_sponsors.py --emit-seeds /tmp/s.txt   # community-list names not on the list yet (for discover --add-file)

Most postings never say whether they sponsor a visa (measured with scripts/measure_visa_wording.py: 85-100% of engineering
postings in the UK, Germany, the Netherlands, Ireland, the UAE and Singapore say nothing), so this records what the employer's
public record says instead:
  uk     GOV.UK register of licensed sponsors: workers (Skilled Worker, Global Business Mobility, Scale-up routes)
  nl     IND public register "Work": recognised sponsors for regular labour and highly skilled migrants
  ie     DETE "employment permits issued to companies", this year and last (permits actually granted)
  us     USCIS H-1B Employer Data Hub: H-1B petitions approved per employer in a fiscal year (T143). uscis.gov answers scripts
         with HTTP 403, so the download is JohnBroberg/H1B_Hub's copy of the hub's own export (FY2024, the newest there), or a
         newer export saved from the hub in a browser and passed with --us-h1b FILE
  ca     ESDC positive LMIA employers list, the last four quarters (Open Government Licence - Canada; T143)
  lists  community lists of companies that hire from India with visa and relocation (abhi-jha's gist, shubheksha's repo) and
         relocate.me's hand-checked list of companies hiring internationally (CC0-1.0; T141), each entry with its places
Germany, the UAE, Singapore, Australia and New Zealand publish no register to download: the rule of each is in job_hunt/visa.py
COUNTRY_RULE (T141, T143), and the lists name companies there. A register says the employer can sponsor, not that a given job does.

A match is "exact" when the names agree once legal suffixes and country words are dropped ("Google (UK) Limited" is Google),
"squashed" when they agree once the spaces are gone too ("Cloudimperiumgames" is "Cloud Imperium Games Ltd"; six letters or
more), "loose" when they agree only once generic words are dropped too ("Atlassian (UK) Operations Limited" is Atlassian). Only
companies on companies.json are written, with the source's own name, so every match can be checked; the registers themselves
are not copied. Run from the repository folder; the downloads go to data/visa_sources/ (git-ignored).
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import html as htmllib
import json
import re
import sys
import unicodedata
from functools import lru_cache
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

ROOT = Path(__file__).resolve().parent.parent
UA = {"User-Agent": "Mozilla/5.0 (RoleSight visa-sponsor builder)"}

UK_CONTENT_API = "https://www.gov.uk/api/content/government/publications/register-of-licensed-sponsors-workers"
UK_PAGE = "https://www.gov.uk/government/publications/register-of-licensed-sponsors-workers"
NL_PAGE = "https://ind.nl/en/public-register-recognised-sponsors/public-register-regular-labour-and-highly-skilled-migrants"
IE_STATS = "https://enterprise.gov.ie/en/what-we-do/workplace-and-skills/employment-permits/statistics/"
GIST_API = "https://api.github.com/gists/5225e0dcc912174dda6b7f01e8a24156"
GIST_PAGE = "https://gist.github.com/abhi-jha/5225e0dcc912174dda6b7f01e8a24156"
SHUBHEKSHA_RAW = "https://raw.githubusercontent.com/shubheksha/companies-sponsoring-visas/master/README.md"
SHUBHEKSHA_PAGE = "https://github.com/shubheksha/companies-sponsoring-visas"
RELOCATE_RAW = "https://raw.githubusercontent.com/AndrewStetsenko/tech-jobs-with-relocation/main/README.md"
RELOCATE_PAGE = "https://github.com/AndrewStetsenko/tech-jobs-with-relocation"   # relocate.me's list, CC0-1.0 (T141)
US_HUB = "https://www.uscis.gov/tools/reports-and-studies/h-1b-employer-data-hub"
US_MIRROR = "https://raw.githubusercontent.com/JohnBroberg/H1B_Hub/main/data/Employer_Information_{year}.csv"
US_MIRROR_PAGE = "https://github.com/JohnBroberg/H1B_Hub"
CA_API = "https://open.canada.ca/data/api/action/package_show?id=90fed587-1364-4f33-a9ee-208181dc0b97"
CA_PAGE = "https://open.canada.ca/data/en/dataset/90fed587-1364-4f33-a9ee-208181dc0b97"
CA_QUARTERS = 4

UK_ROUTES = ("Skilled Worker", "Global Business Mobility: Senior or Specialist Worker", "Scale-up",
             "Global Business Mobility: UK Expansion Worker", "Global Business Mobility: Graduate Trainee")

# --- names ---------------------------------------------------------------------------------------------------------

LEGAL = {"limited", "ltd", "llp", "lp", "plc", "inc", "incorporated", "corp", "corporation", "co", "company", "llc", "gmbh",
         "ag", "se", "bv", "nv", "sa", "sas", "sarl", "srl", "spa", "ab", "as", "asa", "oy", "oyj", "pte", "pty", "pvt",
         "private", "kg", "dac", "uc", "ulc", "unlimited", "branch", "the", "and",
         "pbc", "opco", "pc", "pllc"}           # US: public benefit corp ("ANTHROPIC PBC"), operating company ("OPENAI OPCO LLC")
PLACE = {"uk", "gb", "britain", "great", "england", "scotland", "wales", "ireland", "irl", "nederland", "netherlands",
         "holland", "europe", "european", "emea", "eu", "india", "germany", "deutschland", "us", "usa", "america", "singapore",
         "asia", "apac", "uae", "dubai", "canada"}   # no city names: "Brave London Ltd" is not Brave; "Google Canada" is Google
GENERIC = {"services", "service", "operations", "technologies", "technology", "tech", "solutions", "software", "systems",
           "labs", "lab", "digital", "global", "international", "holdings", "holding", "group", "payments", "bank",
           "platforms", "platform", "ventures", "enterprises", "worldwide", "hq"}


def _tokens(name: str) -> List[str]:
    t = unicodedata.normalize("NFKD", name or "").encode("ascii", "ignore").decode().lower()
    t = re.sub(r"\.(com|io|ai|co)\b", "", t)            # "Booking.com B.V." and "Booking.com" are both "booking"
    t = t.replace("&", " and ").replace(".", "").replace("'", "")
    t = re.sub(r"\b([a-z]) (?=[a-z]\b)", r"\1", re.sub(r"[^a-z0-9]+", " ", t))   # the H-1B file's "L L C", "D B A": "llc", "dba"
    words = t.split()
    return words[words.index("dba") + 1:] if "dba" in words[1:-1] else words   # "X D B A Y": the trade name Y


def exact_key(name: str) -> str:
    """The name without legal suffixes and country words: "Google (UK) Limited" -> "google"; a branch's place goes too
    ("DBS Bank Ltd, London Branch" -> "dbs bank")."""
    t = _tokens(name)
    t = [w for i, w in enumerate(t) if not (i + 1 < len(t) and t[i + 1] == "branch")]
    return " ".join(w for w in t if w not in LEGAL and w not in PLACE)


def squashed_key(name: str) -> str:
    """exact_key with the spaces taken out: a name built from a web address ("Cloudimperiumgames", "Basicfit") against the
    register's ("Cloud Imperium Games Ltd", "Basic-Fit B.V."). Only for names of six letters or more (T144)."""
    key = exact_key(name).replace(" ", "")
    return key if len(key) >= 6 else ""


def loose_key(name: str) -> str:
    """exact_key without generic business words too: "Atlassian (UK) Operations Limited" -> "atlassian"."""
    return " ".join(w for w in exact_key(name).split() if w not in GENERIC)


class Index:
    """Source names by exact and loose key, for matching company names against one source."""

    def __init__(self, entries: Dict[str, Any]):
        self.entries = entries                        # source's own name -> detail (routes, permits, list section)
        self.exact: Dict[str, List[str]] = {}
        self.squashed: Dict[str, List[str]] = {}
        self.loose: Dict[str, List[str]] = {}
        for n in entries:
            if exact_key(n):
                self.exact.setdefault(exact_key(n), []).append(n)
            if squashed_key(n):
                self.squashed.setdefault(squashed_key(n), []).append(n)
            if loose_key(n):
                self.loose.setdefault(loose_key(n), []).append(n)

    def match(self, company: str) -> Tuple[Optional[str], List[str]]:
        """("exact" | "squashed" | "loose" | None, the source names that match)."""
        e, lk = exact_key(company), loose_key(company)
        if e and e in self.exact:
            return "exact", sorted(self.exact[e])
        sq = squashed_key(company)
        if sq and sq in self.squashed:
            return "squashed", sorted(self.squashed[sq])
        if lk and lk in self.loose:
            return "loose", sorted(self.loose[lk])
        return None, []


# --- parsers (each takes a downloaded file) -----------------------------------------------------------------------

def parse_uk(path: Path) -> Dict[str, List[str]]:
    """Organisation -> the worker routes it holds a licence for (only UK_ROUTES)."""
    out: Dict[str, set] = {}
    with open(path, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            route = (r.get("Route") or "").strip()
            name = (r.get("Organisation Name") or "").strip()
            if name and route in UK_ROUTES:
                out.setdefault(name, set()).add(route)
    return {n: sorted(v) for n, v in out.items()}


def parse_nl(text: str) -> Dict[str, str]:
    """Organisation -> KvK (Chamber of Commerce) number, from the register page's table (name in a row header cell)."""
    out: Dict[str, str] = {}
    for tr in re.findall(r"(?s)<tr.*?</tr>", text):
        cells = [htmllib.unescape(re.sub(r"<[^>]+>", "", c)).strip() for c in re.findall(r"(?s)<t[dh][\s>].*?</t[dh]>", tr)]
        if len(cells) >= 2 and cells[0] and re.fullmatch(r"\d{8}", cells[1]):    # skips the header row
            out[cells[0].replace('""', '"')] = cells[1]
    return out


def parse_ie(path: Path) -> Dict[str, int]:
    """Employer -> permits issued in the file's year (its total column)."""
    import openpyxl
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
    out: Dict[str, int] = {}
    total_col = None
    for row in ws.iter_rows(values_only=True):
        if total_col is None:
            hits = [i for i, v in enumerate(row) if isinstance(v, str) and "total" in v.lower()]
            if hits:
                total_col = hits[-1]
            continue
        name = row[0].strip() if isinstance(row[0], str) else ""
        if not name or name.lower() in ("total", "grand total"):
            continue
        n = row[total_col] if total_col < len(row) else None
        out[name] = out.get(name, 0) + (int(n) if isinstance(n, (int, float)) else 0)
    return out


def _int(text: str) -> int:
    return int(text.replace(",", "").strip() or 0)        # the hub writes "1,305"


def parse_us(path: Path) -> Dict[str, Dict[str, List[int]]]:
    """Employer -> {fiscal year: [initial approvals, continuing approvals]}, from an H-1B Employer Data Hub export (tab
    separated, UTF-16 as the hub writes it, or UTF-8); one employer's rows (one per tax id and city) are added up. "Initial"
    is new employment (a cap petition, or a worker changing employer); "continuing" is an extension for a worker already there."""
    raw = path.read_bytes()
    text = raw.decode("utf-16") if raw[:2] in (b"\xff\xfe", b"\xfe\xff") else raw.decode("utf-8-sig")
    rows = csv.reader(text.splitlines(), delimiter="\t")
    head = [h.strip() for h in next(rows, [])]
    col = {h: i for i, h in enumerate(head)}
    out: Dict[str, Dict[str, List[int]]] = {}
    for r in rows:
        if len(r) < len(head):
            continue
        name, year = r[col["Employer (Petitioner) Name"]].strip(), r[col["Fiscal Year"]].strip()
        if not name:
            continue
        counts = out.setdefault(name, {}).setdefault(year, [0, 0])
        counts[0] += _int(r[col["Initial Approval"]])
        counts[1] += _int(r[col["Continuing Approval"]])
    return {n: years for n, years in out.items() if any(sum(c) for c in years.values())}     # denials only: no record


# NOC 2021 occupations of software work (21210-21234 computer and information systems, 21311 computer engineers, 20012 their
# managers): an LMIA for one says more about a software job than one for a cook.
_TECH_NOC = re.compile(r"^(212\d\d|21311|20012)\b")


def parse_ca(path: Path) -> Dict[str, Dict[str, Any]]:
    """Employer -> {"positions", "tech" (positions in software occupations), "streams"} for one quarter's positive LMIA list."""
    import openpyxl
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
    out: Dict[str, Dict[str, Any]] = {}
    col: Dict[str, int] = {}
    for row in ws.iter_rows(values_only=True):
        cells = [str(v).strip() if v is not None else "" for v in row]
        if not col:
            if "Employer" in cells and "Approved Positions" in cells:
                col = {h: i for i, h in enumerate(cells)}
            continue
        name = cells[col["Employer"]] if col["Employer"] < len(cells) else ""
        n = row[col["Approved Positions"]] if col["Approved Positions"] < len(row) else None
        if not name or not isinstance(n, (int, float)):          # the notes under the table
            continue
        e = out.setdefault(name, {"positions": 0, "tech": 0, "streams": []})
        e["positions"] += int(n)
        e["tech"] += int(n) if _TECH_NOC.match(cells[col["Occupation"]]) else 0
        stream = cells[col["Program Stream"]]
        if stream and stream not in e["streams"]:
            e["streams"].append(stream)
    return out


def parse_gist(text: str) -> Dict[str, List[str]]:
    """Company -> the city headings it is listed under ("**London**" then "* [Google](...)")."""
    md = next(iter(json.loads(text)["files"].values()))["content"]
    out: Dict[str, List[str]] = {}
    section = ""
    for line in md.splitlines():
        h = re.fullmatch(r"\s*\*\*(.+?)\*\*\s*", line)
        if h:
            section = h.group(1).strip()
            continue
        m = re.match(r"\s*\*\s*\[([^\]]+)\]", line)
        if m and section:
            out.setdefault(m.group(1).strip(), [])
            if section not in out[m.group(1).strip()]:
                out[m.group(1).strip()].append(section)
    return out


def parse_shubheksha(text: str) -> Dict[str, List[str]]:
    """Company -> its "Location(s)" cell, from the README's table."""
    out: Dict[str, List[str]] = {}
    for line in text.splitlines():
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) < 2 or not cells[0] or cells[0].lower() == "name" or set(cells[0]) <= set("-: "):
            continue
        out[cells[0]] = [re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", p).strip() for p in re.split(r"<br\s*/?>", cells[1]) if p.strip()]
    return out


def parse_relocate(text: str) -> Dict[str, List[str]]:
    """Company -> its places, from the "Companies hiring internationally" table of relocate.me's README (flags as countries:
    "Hamburg 🇩🇪" -> "Hamburg, Germany")."""
    section = text.split("Companies hiring internationally", 1)[-1].split("\n# ", 1)[0]
    out: Dict[str, List[str]] = {}
    for line in section.splitlines():
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) < 2 or not cells[0] or cells[0].lower() == "company name" or set(cells[0]) <= set("-: "):
            continue
        out[cells[0]] = [_flags_to_names(p).strip() for p in re.split(r"<br\s*/?>", cells[1]) if p.strip()]
    return out


_FLAG_NAMES = {"DE": "Germany", "NL": "Netherlands", "GB": "United Kingdom", "IE": "Ireland", "AE": "United Arab Emirates",
               "SG": "Singapore", "IN": "India"}
_COUNTRIES = ROOT / "job_hunt" / "data" / "countries.json"


@lru_cache(maxsize=1)
def _country_names() -> Dict[str, str]:
    """ISO2 code -> country name: the names above, else the place data the scanner reads (T142), so every flag gets one (T143)."""
    try:
        names = {code: c["name"] for code, c in json.loads(_COUNTRIES.read_text(encoding="utf-8"))["countries"].items()}
    except (OSError, ValueError, KeyError):
        names = {}
    return {**names, **_FLAG_NAMES}


def _flags_to_names(text: str) -> str:
    names = _country_names()

    def name(m: "re.Match[str]") -> str:
        code = "".join(chr(ord(ch) - 0x1F1E6 + 65) for ch in m.group(1))
        return ", " + names.get(code, code)
    return re.sub(r"\s*([\U0001F1E6-\U0001F1FF]{2})", name, text).replace("🌍", "").strip(" ,")


# --- downloads -----------------------------------------------------------------------------------------------------

def _get(url: str, timeout: int = 120):
    import requests
    r = requests.get(url, headers=UA, timeout=timeout)
    r.raise_for_status()
    return r


def _download_us(dest: Path, us_file: Optional[Path]) -> Dict[str, Any]:
    """The H-1B hub export: `us_file` when given (saved from the hub in a browser), else the newest year of the mirror."""
    meta = {"title": "USCIS H-1B Employer Data Hub (H-1B petitions approved per employer)", "page": US_HUB,
            "licence": "US government work (public domain)", "files": ["us_h1b.csv"]}
    if us_file:
        (dest / "us_h1b.csv").write_bytes(us_file.read_bytes())
        meta["file"] = str(us_file.name)
    else:
        for year in range(dt.date.today().year, 2020, -1):          # the newest year the copy has (2024 on 2026-10-21)
            try:
                content = _get(US_MIRROR.format(year=year), timeout=300).content
            except Exception:                                       # 404: not there yet
                continue
            (dest / "us_h1b.csv").write_bytes(content)
            meta.update(file=US_MIRROR.format(year=year), copy=US_MIRROR_PAGE + " (MIT; the hub's export, unchanged)")
            break
        else:
            raise RuntimeError("no H-1B hub export found on the mirror; save one from the hub and pass --us-h1b FILE")
    meta["years"] = sorted({y for counts in parse_us(dest / "us_h1b.csv").values() for y in counts})
    return meta


def _download_ca(dest: Path) -> Dict[str, Any]:
    """The last CA_QUARTERS quarters of the positive LMIA employers list (English xlsx files, found with the CKAN API)."""
    res = _get(CA_API).json()["result"]["resources"]
    quarters = sorted(((m.group(1), r["url"]) for r in res if r.get("format", "").upper().startswith("XLS")
                       and "en" in (r.get("language") or ["en"]) and (m := re.match(r"(\d{4}Q[1-4])\b", r.get("name", "")))))
    files = []
    for q, url in quarters[-CA_QUARTERS:]:
        (dest / f"ca_lmia_{q}.xlsx").write_bytes(_get(url).content)
        files.append(f"ca_lmia_{q}.xlsx")
    return {"title": "ESDC Temporary Foreign Worker Program: positive LMIA employers list", "page": CA_PAGE,
            "quarters": [q for q, _ in quarters[-CA_QUARTERS:]], "licence": "Open Government Licence - Canada",
            "files": files}


def download(dest: Path, us_file: Optional[Path] = None) -> Dict[str, Dict[str, Any]]:
    """Fetch every source into `dest`; returns what each one is (url, date, licence, file)."""
    dest.mkdir(parents=True, exist_ok=True)
    meta: Dict[str, Dict[str, Any]] = {}

    c = _get(UK_CONTENT_API).json()
    csv_url = next(a["url"] for a in c["details"]["attachments"] if a.get("content_type") == "text/csv")
    (dest / "uk_register.csv").write_bytes(_get(csv_url).content)
    meta["uk"] = {"title": c.get("title", "Register of licensed sponsors: workers"), "page": UK_PAGE, "file": csv_url,
                  "published": (c.get("public_updated_at") or "")[:10], "licence": "Open Government Licence v3.0",
                  "files": ["uk_register.csv"]}

    page = _get(NL_PAGE).text
    (dest / "nl_register.html").write_text(page, encoding="utf-8")
    m = re.search(r"last updated on (\d{1,2} \w+ \d{4})", page)
    meta["nl"] = {"title": "IND public register Work (recognised sponsors)", "page": NL_PAGE,
                  "published": dt.datetime.strptime(m.group(1), "%d %B %Y").date().isoformat() if m else "",
                  "licence": "none stated on the page", "files": ["nl_register.html"]}

    stats = _get(IE_STATS).text
    years = sorted({int(y) for y in re.findall(r"(?i)/en/publications/employment-permit-statistics-(\d{4})\.html", stats)})[-2:]
    files = []
    for y in years:
        year_page = _get(f"https://enterprise.gov.ie/en/publications/employment-permit-statistics-{y}.html").text
        href = next(h for h in re.findall(r'href="([^"]+\.xlsx)"', year_page) if "issued-to-companies" in h)
        name = f"ie_companies_{y}.xlsx"
        (dest / name).write_bytes(_get(href if href.startswith("http") else "https://enterprise.gov.ie" + href).content)
        files.append(name)
    meta["ie"] = {"title": "DETE employment permits issued to companies", "page": IE_STATS, "years": years,
                  "licence": "Government of Ireland copyright: reuse with the source and copyright acknowledged",
                  "files": files}

    meta["us"] = _download_us(dest, us_file)
    meta["ca"] = _download_ca(dest)

    (dest / "gist.json").write_text(_get(GIST_API).text, encoding="utf-8")
    (dest / "shubheksha.md").write_text(_get(SHUBHEKSHA_RAW).text, encoding="utf-8")
    (dest / "relocate.md").write_text(_get(RELOCATE_RAW).text, encoding="utf-8")
    meta["lists"] = {"title": "Community lists of companies hiring from India / internationally with visa and relocation",
                     "pages": [GIST_PAGE, SHUBHEKSHA_PAGE, RELOCATE_PAGE],
                     "licence": "gist and shubheksha: none stated (only matched names are recorded); relocate.me: CC0-1.0",
                     "files": ["gist.json", "shubheksha.md", "relocate.md"]}
    for k in meta:
        meta[k]["downloaded"] = dt.date.today().isoformat()
    (dest / "sources.json").write_text(json.dumps(meta, indent=2), encoding="utf-8")
    return meta


# --- build ---------------------------------------------------------------------------------------------------------

def load_sources(src: Path) -> Dict[str, Index]:
    meta = json.loads((src / "sources.json").read_text(encoding="utf-8"))
    ie: Dict[str, Dict[str, int]] = {}
    for name, year in zip(meta["ie"]["files"], meta["ie"]["years"]):
        for emp, n in parse_ie(src / name).items():
            ie.setdefault(emp, {})[str(year)] = n
    lists: Dict[str, List[str]] = {}
    parsed_lists = [("abhi-jha gist", parse_gist((src / "gist.json").read_text(encoding="utf-8"))),
                    ("shubheksha list", parse_shubheksha((src / "shubheksha.md").read_text(encoding="utf-8")))]
    if (src / "relocate.md").exists():            # added in T141: an older download has no relocate.me file
        parsed_lists.append(("relocate.me list", parse_relocate((src / "relocate.md").read_text(encoding="utf-8"))))
    for label, parsed in parsed_lists:
        for company, where in parsed.items():
            lists.setdefault(company, []).extend(f"{label}: {w}" for w in where or [""])
    out = {"uk": Index(parse_uk(src / "uk_register.csv")),
           "nl": Index(parse_nl((src / "nl_register.html").read_text(encoding="utf-8"))), "ie": Index(ie)}
    if "us" in meta:                                # added in T143: an older download has neither
        out["us"] = Index(parse_us(src / "us_h1b.csv"))
    if "ca" in meta:
        ca: Dict[str, Dict[str, Any]] = {}
        for name, quarter in zip(meta["ca"]["files"], meta["ca"]["quarters"]):
            for emp, e in parse_ca(src / name).items():
                ca.setdefault(emp, {})[quarter] = e
        out["ca"] = Index(ca)
    out["lists"] = Index(lists)
    return out


Candidate = Tuple[str, str, str, str, Any]      # company, source, "exact" | "squashed" | "loose", the source's name, its detail


def load_review(path: Path) -> Dict[str, Dict[str, List[str]]]:
    """Verdicts read by hand: {"accepted": {company: [source names]}, "rejected": {...}}; no file means none."""
    r = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    return {"accepted": r.get("accepted", {}), "rejected": r.get("rejected", {})}


def status(company: str, source_name: str, match: str, review: Dict[str, Dict[str, List[str]]]) -> str:
    """"kept" (an exact match, or a loose one accepted by hand), "rejected" by hand, or "unreviewed" (a loose one)."""
    if source_name in review["rejected"].get(company, []):
        return "rejected"
    if match in ("exact", "squashed") or source_name in review["accepted"].get(company, []):
        return "kept"
    return "unreviewed"


def candidates(companies: Iterable[Dict[str, Any]], indexes: Dict[str, Index]) -> List[Candidate]:
    """Every match the names allow, before the review."""
    out: List[Candidate] = []
    for c in companies:
        for source, idx in indexes.items():
            kind, names = idx.match(c["name"])
            out.extend((c["name"], source, kind, n, idx.entries[n]) for n in names if kind)
    return out


def match_companies(cands: List[Candidate], review: Dict[str, Dict[str, List[str]]]) -> Dict[str, List[Dict[str, Any]]]:
    """Company name -> evidence: [{"source", "match", "names": {source's name: detail}}], kept matches only."""
    out: Dict[str, List[Dict[str, Any]]] = {}
    for company, source, kind, name, detail in cands:
        if status(company, name, kind, review) != "kept":
            continue
        ev = out.setdefault(company, [])
        e = next((x for x in ev if x["source"] == source), None)
        if e is None:
            e = {"source": source, "match": kind, "names": {}}
            ev.append(e)
        e["names"][name] = detail
    return out


def build(src: Path, companies: List[Dict[str, Any]], review: Dict[str, Dict[str, List[str]]]
          ) -> Tuple[Dict[str, Any], List[Candidate]]:
    meta = json.loads((src / "sources.json").read_text(encoding="utf-8"))
    indexes = load_sources(src)
    for k, idx in indexes.items():
        meta[k]["entries"] = len(idx.entries)
        meta[k].pop("files", None)
    cands = candidates(companies, indexes)
    return ({"about": "Companies on companies.json that appear in public visa-sponsor sources (T133). A register means the "
                      "employer can sponsor a visa, not that a given job does. Each match names the source's own entry; "
                      "loose name matches are kept only when accepted by hand in scripts/visa_sponsor_review.json. "
                      "Rebuild: python scripts/build_visa_sponsors.py",
             "built": dt.date.today().isoformat(), "sources": meta, "companies": match_companies(cands, review)}, cands)


def seeds(src: Path, companies: List[Dict[str, Any]]) -> List[str]:
    """Community-list names that match no company on the list (they may be worth a discover run)."""
    have = Index({c["name"]: "" for c in companies})
    return sorted(n for n in load_sources(src)["lists"].entries if not have.match(n)[0])


def write_report(path: str, cands: List[Candidate], review: Dict[str, Dict[str, List[str]]],
                 companies: List[Dict[str, Any]]) -> int:
    ats = {c["name"]: c.get("ats", "") for c in companies}
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Status", "Company", "ATS", "Source", "Match", "Source's name", "Detail"])
        for company, source, kind, name, detail in sorted(cands, key=lambda c: (c[0].lower(), c[1], c[3])):
            w.writerow([status(company, name, kind, review), company, ats.get(company, ""), source, kind, name,
                        json.dumps(detail, ensure_ascii=False)])
    return len(cands)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--companies", default="companies.json")
    ap.add_argument("--out", default="job_hunt/data/visa_sponsors.json")
    ap.add_argument("--sources", default="data/visa_sources", help="where the downloads go (default data/visa_sources)")
    ap.add_argument("--offline", action="store_true", help="use the files downloaded last time")
    ap.add_argument("--us-h1b", help="an H-1B Employer Data Hub export saved in a browser (uscis.gov refuses scripts); "
                                     "default: the newest year on the JohnBroberg/H1B_Hub copy")
    ap.add_argument("--review", default=str(ROOT / "scripts" / "visa_sponsor_review.json"),
                    help="matches accepted / rejected by hand (default scripts/visa_sponsor_review.json)")
    ap.add_argument("--report", help="CSV of every match with its status (kept, rejected, unreviewed), to read by hand")
    ap.add_argument("--emit-seeds", help="community-list names not on the list, one per line (for discover --add-file)")
    args = ap.parse_args(argv)
    src = Path(args.sources)
    if args.offline:
        if not (src / "sources.json").exists():
            sys.exit(f"build_visa_sponsors: nothing downloaded in {src} yet; run once without --offline")
        if args.us_h1b:                             # a newer hub export replaces the stored one; nothing else is fetched
            meta = json.loads((src / "sources.json").read_text(encoding="utf-8"))
            meta["us"] = dict(_download_us(src, Path(args.us_h1b)), downloaded=dt.date.today().isoformat())
            (src / "sources.json").write_text(json.dumps(meta, indent=2), encoding="utf-8")
    else:
        download(src, Path(args.us_h1b) if args.us_h1b else None)
    companies = json.loads(Path(args.companies).read_text(encoding="utf-8"))
    review = load_review(Path(args.review))
    data, cands = build(src, companies, review)
    Path(args.out).write_text(json.dumps(data, indent=1, ensure_ascii=False, sort_keys=True) + "\n", encoding="utf-8")
    by = {k: sum(any(e["source"] == k for e in ev) for ev in data["companies"].values()) for k in data["sources"]}
    print(f"{len(data['companies'])} of {len(companies)} companies have evidence ({', '.join(f'{k} {v}' for k, v in by.items())}); "
          f"wrote {args.out}")
    counts = {s: sum(status(c, n, k, review) == s for c, _, k, n, _ in cands) for s in ("kept", "rejected", "unreviewed")}
    print(f"matches: {counts['kept']} kept, {counts['rejected']} rejected by hand, {counts['unreviewed']} loose ones unreviewed"
          + (" (not used: read them with --report, then accept or reject them in the review file)" if counts["unreviewed"] else ""))
    if args.report:
        print(f"{write_report(args.report, cands, review, companies)} matches in {args.report}")
    if args.emit_seeds:
        names = seeds(src, companies)
        Path(args.emit_seeds).write_text("".join(f"{n}\n" for n in names), encoding="utf-8")
        print(f"{len(names)} community-list names not on the list in {args.emit_seeds}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
